"""Build the editable report and slide deck from the project's saved results."""

from pathlib import Path
import re

from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
SLIDES = ROOT / "presentation"
SLIDES.mkdir(exist_ok=True)


def clean(text):
    return text.replace("**", "").replace("`", "")


def build_report():
    source = (REPORT / "Final_Project_Report.md").read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = DocInches(0.8)
    section.bottom_margin = DocInches(0.7)
    section.left_margin = section.right_margin = DocInches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = DocPt(10)
    normal.paragraph_format.space_after = DocPt(6)
    for name, size in [("Title", 21), ("Heading 1", 13)]:
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = DocPt(size)
        style.font.bold = True
        style.font.color.rgb = DocColor(0, 0, 0)
        style.paragraph_format.space_before = DocPt(12)
        style.paragraph_format.space_after = DocPt(5)

    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(clean(line[2:]), "Title")
        elif line.startswith("## "):
            doc.add_paragraph(clean(line[3:]), "Heading 1")
        elif line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].startswith("| "):
                cells = [clean(x.strip()) for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", x) for x in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
            i -= 1
        else:
            doc.add_paragraph(clean(line))
        i += 1
    footer = section.footer.paragraphs[0]
    footer.text = "NNDL Deepfake Detection Project"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = "Caption"
    output = REPORT / "Final_Project_Report.docx"
    doc.save(output)
    return output


def textbox(slide, text, x, y, w, h, size=20, bold=False, color=(25, 35, 52)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for n, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(10)
    return box


def title(slide, heading, subtitle=None):
    textbox(slide, heading, 0.7, 0.45, 11.8, 0.7, 31, True)
    if subtitle:
        textbox(slide, subtitle, 0.72, 1.28, 11.5, 0.55, 16, False, (84, 97, 113))


def picture(slide, path, x, y, w=None, h=None):
    kwargs = {"left": Inches(x), "top": Inches(y)}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    slide.shapes.add_picture(str(ROOT / path), **kwargs)


def build_slides():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    textbox(s, "Deepfake image detection", 0.75, 1.45, 11.8, 1.1, 43, True)
    textbox(s, "Custom CNN versus EfficientNet-B0", 0.78, 2.7, 11.5, 0.7, 28)
    textbox(s, "NNDL project  |  Final results and demonstration", 0.8, 5.8, 11.5, 0.4, 16, False, (84, 97, 113))

    s = prs.slides.add_slide(blank)
    title(s, "Task and dataset")
    textbox(s, "Binary classification of face images: FAKE or REAL", 0.75, 1.65, 11.8, 0.55, 25, True)
    textbox(s, "Training: 8,000 images  (4,000 per class)\nValidation: 1,000 images  (500 per class)\nHeld-out test: 1,000 images  (500 per class)", 0.78, 2.55, 11.7, 2.2, 22)
    textbox(s, "The integrity report found no corrupt images or exact duplicates crossing splits.", 0.78, 5.55, 11.7, 0.8, 19, False, (84, 97, 113))

    s = prs.slides.add_slide(blank)
    title(s, "Two models, one test set")
    textbox(s, "Custom CNN", 0.8, 1.55, 5.1, 0.5, 25, True)
    textbox(s, "Trained from scratch\n224 × 224 input\nAdamW and early stopping\n11 logged epochs", 0.8, 2.25, 5.4, 3.4, 21)
    textbox(s, "EfficientNet-B0", 6.8, 1.55, 5.5, 0.5, 25, True)
    textbox(s, "ImageNet transfer learning\nClassifier head, then feature fine-tuning\nBest validation-loss checkpoint\n25 logged epochs", 6.8, 2.25, 5.7, 3.4, 21)

    s = prs.slides.add_slide(blank)
    title(s, "EfficientNet leads on every metric", "Held-out test set; FAKE is the positive class")
    picture(s, "results/comparison/cnn_vs_efficientnet_metrics.png", 2.05, 1.8, w=9.2)

    s = prs.slides.add_slide(blank)
    title(s, "96.9% accuracy on the held-out test set")
    textbox(s, "EfficientNet-B0", 0.8, 1.55, 5.3, 0.5, 24, True)
    textbox(s, "Accuracy  96.9%\nPrecision  96.6%\nRecall  97.2%\nF1  96.9%\nROC-AUC  0.9963", 0.8, 2.3, 5.0, 3.6, 23)
    textbox(s, "Custom CNN", 6.7, 1.55, 5.4, 0.5, 24, True)
    textbox(s, "Accuracy  72.3%\nPrecision  69.9%\nRecall  78.4%\nF1  73.9%\nROC-AUC  0.8018", 6.7, 2.3, 5.0, 3.6, 23)

    s = prs.slides.add_slide(blank)
    title(s, "Error analysis", "Rows are true classes; columns are predicted classes")
    picture(s, "results/comparison/cnn_vs_efficientnet_confusion_matrices.png", 0.7, 1.72, w=8.0)
    textbox(s, "CNN: 277 errors\nEfficientNet: 31 errors\n\nREAL false alarms:\n169 → 17\n\nMissed FAKE images:\n108 → 14", 9.05, 1.78, 3.5, 4.7, 20)

    s = prs.slides.add_slide(blank)
    title(s, "Training behavior")
    picture(s, "results/comparison/cnn_vs_efficientnet_accuracy_comparison.png", 0.65, 1.55, w=7.5)
    textbox(s, "CNN final train / validation gap: 6.25 points.\n\nEfficientNet gap at the selected epoch: 1.35 points.\n\nBoth show some late-stage overfitting; evaluation uses the best validation-loss checkpoint.", 8.45, 1.8, 4.1, 4.7, 18)

    s = prs.slides.add_slide(blank)
    title(s, "Local image-upload demo")
    textbox(s, "Upload a face image", 0.8, 1.7, 11.5, 0.5, 28, True)
    textbox(s, "EfficientNet-B0 applies the same test preprocessing and computes P(FAKE).", 0.8, 2.5, 11.6, 0.85, 22)
    textbox(s, "Output: FAKE or REAL  •  predicted-class confidence  •  P(FAKE)  •  P(REAL)", 0.8, 3.75, 11.7, 1.2, 22, True)
    textbox(s, "Run: python -m streamlit run app/demo.py", 0.8, 5.7, 11.6, 0.6, 18, False, (84, 97, 113))

    s = prs.slides.add_slide(blank)
    title(s, "Conclusion and limits")
    textbox(s, "EfficientNet-B0 meets the project targets and improves accuracy by 24.6 percentage points over the CNN.", 0.78, 1.65, 11.6, 1.1, 25, True)
    textbox(s, "The result is measured on one held-out dataset. Performance across other generators or cameras is unknown.\n\nThe demo displays raw model probabilities, which are not calibrated authenticity guarantees.", 0.8, 3.25, 11.5, 2.2, 21)

    output = SLIDES / "Final_Project_Presentation.pptx"
    prs.save(output)
    return output


if __name__ == "__main__":
    print(build_report())
    print(build_slides())
